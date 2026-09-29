import streamlit as st
import json
import random
import time
from pypdf import PdfReader
from pptx import Presentation
from google import genai
from google.genai import types

st.set_page_config(page_title="StudyCard Studio ✨", page_icon="📚", layout="centered")

if "master_deck" not in st.session_state:
    st.session_state["master_deck"] = []
if "quiz_index" not in st.session_state:
    st.session_state["quiz_index"] = 0
if "show_answer" not in st.session_state:
    st.session_state["show_answer"] = False

st.title("📚 StudyCard Studio")
st.caption("Upload your chapters to build your master exam deck, then test yourself in Quiz Mode.")

api_key = st.secrets.get("GEMINI_API_KEY")

tab_create, tab_quiz, tab_export = st.tabs(["➕ Add Chapters", "🎯 Practice Quiz", "📥 Master Export"])

def extract_text(file):
    text = ""
    if file.name.endswith(".pdf"):
        reader = PdfReader(file)
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content + "\n"
    elif file.name.endswith(".pptx"):
        prs = Presentation(file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text += shape.text + "\n"
    return text

# --- TAB 1: ADD CHAPTERS ---
with tab_create:
    st.subheader("Add Study Material")
    uploaded_file = st.file_uploader("Upload Chapter Slides (.pptx) or PDF (.pdf)", type=["pdf", "pptx"])
    pasted_notes = st.text_area("Or paste chapter notes:", height=100, placeholder="Paste notes here...")
    
    col1, col2 = st.columns(2)
    with col1:
        num_cards = st.slider("Cards to generate for this chapter:", min_value=5, max_value=35, value=15)
    with col2:
        chapter_tag = st.text_input("Chapter Name / Tag:", placeholder="e.g., Chapter 1")

    if st.button("✨ Generate & Add to Master Deck", type="primary", use_container_width=True):
        if not api_key:
            st.error("API Key missing. Please check your Streamlit Secrets.")
        else:
            combined_text = ""
            if uploaded_file:
                combined_text += extract_text(uploaded_file)
            if pasted_notes:
                combined_text += "\n" + pasted_notes

            if len(combined_text.strip()) < 30:
                st.warning("Please upload a file or enter readable notes.")
            else:
                with st.spinner("Analyzing chapter and crafting exam-style questions..."):
                    safe_material = combined_text[:40000]
                    prompt = f"""
                    You are an expert exam tutor. Create exactly {num_cards} high-yield exam-prep flashcards based on this material.
                    Focus on application, distinguishing easily confused concepts, and cause/effect mechanisms.
                    Return JSON only in this exact format:
                    [
                      {{"question": "Exam-style question here", "answer": "Clear, direct explanation here"}}
                    ]

                    Material:
                    {safe_material}
                    """

                    client = genai.Client(api_key=api_key.strip())
                    res = None
                    last_error = ""

                    # Auto-retries up to 3 times with a short pause if traffic spikes
                    for attempt in range(1, 4):
                        try:
                            res = client.models.generate_content(
                                model="gemini-2.5-flash",
                                contents=prompt,
                                config=types.GenerateContentConfig(response_mime_type="application/json")
                            )
                            if res and res.text:
                                break
                        except Exception as e:
                            last_error = str(e)
                            if "503" in str(e) or "UNAVAILABLE" in str(e):
                                time.sleep(3)
                                continue
                            else:
                                break

                    if res and res.text:
                        try:
                            new_cards = json.loads(res.text)
                            tag = chapter_tag.strip() if chapter_tag.strip() else f"Ch {len(st.session_state['master_deck'])//10 + 1}"
                            for card in new_cards:
                                card["tag"] = tag

                            st.session_state["master_deck"].extend(new_cards)
                            st.session_state["quiz_index"] = 0
                            st.session_state["show_answer"] = False
                            st.success(f"Added {len(new_cards)} cards! Master Deck now has {len(st.session_state['master_deck'])} total cards.")
                        except Exception as e:
                            st.error(f"Error parsing flashcards: {e}")
                    else:
                        st.error(f"Generation error: {last_error}")

    if st.session_state["master_deck"]:
        st.divider()
        col_info, col_clear = st.columns([3, 1])
        with col_info:
            st.info(f"**Current Master Deck:** {len(st.session_state['master_deck'])} cards ready for review.")
        with col_clear:
            if st.button("🗑️ Reset Deck", help="Clear all cards and start a new exam deck"):
                st.session_state["master_deck"] = []
                st.session_state["quiz_index"] = 0
                st.rerun()

# --- TAB 2: PRACTICE QUIZ ---
with tab_quiz:
    deck = st.session_state["master_deck"]
    if not deck:
        st.info("No cards in the deck yet. Upload a chapter in the first tab to begin!")
    else:
        st.subheader("🎯 Cumulative Exam Quiz")
        
        c1, c2, c3 = st.columns([1, 2, 1])
        with c1:
            if st.button("🔀 Shuffle Deck"):
                random.shuffle(st.session_state["master_deck"])
                st.session_state["quiz_index"] = 0
                st.session_state["show_answer"] = False
                st.rerun()
        with c2:
            st.write(f"Card **{st.session_state['quiz_index'] + 1}** of **{len(deck)}**")
        with c3:
            current_tag = deck[st.session_state['quiz_index']].get("tag", "")
            if current_tag:
                st.caption(f"🏷️ {current_tag}")

        current_card = deck[st.session_state["quiz_index"]]

        st.markdown(
            f"""
            <div style="background-color: #f7f9fc; border: 2px solid #e1e8ed; border-radius: 12px; padding: 25px; min-height: 180px; margin: 15px 0;">
                <h4 style="color: #1f2937; margin-bottom: 10px;">Question:</h4>
                <p style="font-size: 1.15rem; color: #111827;">{current_card['question']}</p>
            </div>
            """,
            unsafe_allow_html=True
        )

        if st.session_state["show_answer"]:
            st.markdown(
                f"""
                <div style="background-color: #f0fdf4; border: 2px solid #86efac; border-radius: 12px; padding: 25px; margin-bottom: 20px;">
                    <h4 style="color: #15803d; margin-bottom: 10px;">Answer:</h4>
                    <p style="font-size: 1.1rem; color: #166534;">{current_card['answer']}</p>
                </div>
                """,
                unsafe_allow_html=True
            )
            if st.button("🙈 Hide Answer"):
                st.session_state["show_answer"] = False
                st.rerun()
        else:
            if st.button("👀 Show Answer / Flip Card", type="secondary", use_container_width=True):
                st.session_state["show_answer"] = True
                st.rerun()

        col_prev, col_next = st.columns(2)
        with col_prev:
            if st.button("⬅️ Previous Card", use_container_width=True):
                if st.session_state["quiz_index"] > 0:
                    st.session_state["quiz_index"] -= 1
                    st.session_state["show_answer"] = False
                    st.rerun()
        with col_next:
            if st.button("Next Card ➡️", type="primary", use_container_width=True):
                if st.session_state["quiz_index"] < len(deck) - 1:
                    st.session_state["quiz_index"] += 1
                    st.session_state["show_answer"] = False
                    st.rerun()
                else:
                    st.balloons()
                    st.success("🎉 You've completed the entire review deck!")

# --- TAB 3: MASTER EXPORT ---
with tab_export:
    deck = st.session_state["master_deck"]
    if not deck:
        st.info("No cards generated yet.")
    else:
        st.subheader("📥 Export Cumulative Deck")
        st.caption("Copy this block and paste it directly into Quizlet or Anki under 'Import' to keep all chapters together:")
        formatted_export = "\n".join([f"[{c.get('tag', 'Exam')}] {c['question']}\t{c['answer']}" for c in deck])
        st.text_area("Tab-separated text:", value=formatted_export, height=220)


     
             


   
